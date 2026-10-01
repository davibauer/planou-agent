#!/usr/bin/env python3
"""Job application in a few commands, from a CV PROFILE saved in the config (cv_profiles): the CV always comes out of the same
reviewed track (experiences, order, adjusted dates), and only the Summary (and optionally the Skills line) changes per job.

  application.py --profiles                                   # saved CV profiles
  application.py <id> --profile devsecops --summary s.txt [--skills "..."] [--company "Name"] [--force]
      -> builds the DOCX (cv_docx), checks the timeline (order, overlap, gap), generates the HTML (docx_html) and prints
         the next step: create the Google Doc through the Drive connector with that HTML and export the PDF.
  application.py <id> --pdf <tool-result.txt> --gdoc <url>  # after Drive: saves the PDF, checks it against the DOCX and records it

Config (config.json, via criteria.py): "cv_dir": folder of the generated CVs; "cv_profiles": {"<name>": {"base": "<docx>",
"sources": {"A": "<docx>"}, "experiences": [["A", "Role at Company"], ...], "dates": {...}, "languages": [...],
"file_name": "Ana Lima, Resume, {aaaa_mm} ({empresa})"}}. The summary comes from Claude, per job, only with facts from perfil.md.
"""
import os, re, sys, json, argparse, subprocess
from datetime import datetime

HERE = os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0, HERE)
import criteria
from watch_core import fileio

MONTHS = {m: i for i, m in enumerate(['january', 'february', 'march', 'april', 'may', 'june', 'july', 'august', 'september',
                                      'october', 'november', 'december'], 1)}
MONTHS.update({m: i for i, m in enumerate(['janeiro', 'fevereiro', 'março', 'abril', 'maio', 'junho', 'julho', 'agosto',
                                           'setembro', 'outubro', 'novembro', 'dezembro'], 1)})
CURRENT = re.compile(r'current|present|atual|hoje|momento', re.I)
CONCURRENT = re.compile(r'^(co-?)?founder\b|^fundador', re.I)    # own company in parallel does not count as overlap
GAP_MONTHS = 2
CHECK_SINCE = 2016            # a gap before that (old career) does not matter

def _month(txt):
    m = re.search(r'([A-Za-zç]+)\s+(\d{4})', txt or '')
    if not m or m.group(1).lower() not in MONTHS: return None
    return int(m.group(2)) * 12 + MONTHS[m.group(1).lower()] - 1

def period(line, today):
    """'September 2024 - February 2025 | Ohio' -> (start, end) in absolute months; 'September 2024 | Ohio' -> end = start."""
    parts = re.split(r'\s[-–]\s', line.split('|')[0])
    start = _month(parts[0])
    if start is None: return None
    if len(parts) < 2: return start, start
    end = today if CURRENT.search(parts[1]) else _month(parts[1])
    return start, end if end is not None else start

def experiences(docx_path):
    import docx, cv_docx
    ps = list(docx.Document(docx_path).paragraphs); out = []
    for i, p in enumerate(ps):
        if cv_docx.is_experience_title(p):
            date = next((ps[j].text.strip() for j in range(i + 1, min(i + 3, len(ps))) if ps[j].text.strip()), '')
            out.append((p.text.strip().split(' | ')[0], date))
    return out

def timeline(docx_path, parallel=()):
    """Date problems in the generated CV: out of order (most recent first), overlap between experiences (own company
    in parallel is allowed), gap > GAP_MONTHS since CHECK_SINCE, experience with an open/missing end date.
    `parallel`: titles of experiences that run in parallel on purpose (personal project, side work) and stay out
    of the order and overlap check (CV profile: key "parallel")."""
    h = datetime.now(); today = h.year * 12 + h.month - 1
    exps = [(t, d, period(d, today)) for t, d in experiences(docx_path)]
    prob = []
    for t, d, p in exps:
        if p is None: prob.append(f'data ilegivel: {t} ({d})')
        elif p[0] == p[1] and not CURRENT.search(d) and ' - ' not in d and ' – ' not in d: prob.append(f'sem data de fim: {t} ({d})')
    seq = [(t, p) for t, d, p in exps if p and not CONCURRENT.search(t) and not any(t.startswith(x) for x in parallel)]
    for (t1, p1), (t2, p2) in zip(seq, seq[1:]):     # t1 comes first in the CV: it must be the most recent
        if p2[0] > p1[0]: prob.append(f'fora de ordem: "{t2}" comeca depois de "{t1}"')
        # leaving and joining in the same month is a normal transition; overlap is the previous one ending AFTER the month the next one starts
        if p2[1] > p1[0]: prob.append(f'sobreposicao: "{t2}" vai ate depois do inicio de "{t1}"')
        elif p2[1] // 12 >= CHECK_SINCE and p1[0] - p2[1] - 1 > GAP_MONTHS:
            prob.append(f'buraco de {p1[0] - p2[1] - 1} meses entre "{t2}" e "{t1}"')
    return exps, prob

def _state():
    import paths; p = paths.STATE; paths.lock_state(); return p, json.load(open(p, encoding='utf-8'))

def _save(p, s):
    fileio.write_json(p, s, indent=1, ensure_ascii=False)

def assemble(a, cfg):
    profiles = cfg.get('cv_profiles') or {}
    if a.perfil not in profiles: sys.exit(f'perfil "{a.perfil}" nao existe; perfis: {", ".join(profiles) or "nenhum (config cv_profiles)"}')
    pf = dict(profiles[a.perfil])
    sp, s = _state(); v = (s.get('jobs') or {}).get('tracked', {}).get(str(a.id), {})
    company = a.empresa or v.get('company') or sys.exit('empresa desconhecida: passe --company')
    company_file = re.sub(r'[,.]? (inc|ltda|s\.?a|llc|ltd|corp|co)\.?$', '', company, flags=re.I).strip()
    summary = [x.strip() for x in open(a.summary, encoding='utf-8').read().split('\n\n') if x.strip()]
    if any('—' in x for x in summary): sys.exit('summary com travessao: trocar por virgula, dois-pontos ou ponto')
    plan = {k: pf[k] for k in ('base', 'sources', 'experiences', 'dates', 'languages', 'new_entries', 'cuts', 'intros') if k in pf}
    plan['summary'] = summary
    if a.skills: plan['skills'] = a.skills
    name = (pf.get('file_name') or 'Resume, {aaaa_mm} ({empresa})').format(aaaa_mm=datetime.now().strftime('%Y-%m'), empresa=company_file)
    import paths; folder = os.path.expanduser(cfg.get('cv_dir') or paths.CV_DIR); os.makedirs(folder, exist_ok=True)
    docx_p, html_p = os.path.join(folder, name + '.docx'), os.path.join(folder, name + '.html')
    import cv_docx, docx_html
    cv_docx.build(plan, docx_p)
    exps, prob = timeline(docx_p, pf.get('parallel') or ())
    print(f'DOCX: {docx_p}')
    for t, d, _ in exps[:6]: print(f'   {t} | {d}')
    if prob:
        print('CRONOLOGIA COM PROBLEMA:'); [print('   - ' + x) for x in prob]
        if not a.forcar: sys.exit('corrigir o perfil (experiences/dates) ou rodar com --force se estiver certo')
    else: print('cronologia ok: mais recente primeiro, sem sobreposicao, sem buraco')
    open(html_p, 'w', encoding='utf-8').write(docx_html.convert(docx_p))
    report_gaps(a.id, v, docx_p)
    v.update(cv_docx=docx_p, cv_profile=a.perfil); s.setdefault('jobs', {}).setdefault('tracked', {})[str(a.id)] = v; _save(sp, s)
    print(f'HTML: {html_p}')
    print('\nPROXIMO PASSO (conector Google Drive):')
    print(f'  1. create_file(title="{name}", contentMimeType="text/html", textContent=<conteudo de {html_p}>) -> guardar o link do Doc')
    print('  2. download_file_content(fileId=<id do Doc>, exportMimeType="application/pdf") -> o harness salva o resultado num arquivo')
    print(f'  3. python3 {os.path.abspath(__file__)} {a.id} --pdf <arquivo do resultado> --gdoc <link do Doc>')

# Named technologies a screener checks one by one (26/09/2026, Polanov: the posting named Azure Functions, Service Bus,
# Data Factory, PostgreSQL, Cosmos DB and Snowflake, and the CV cited none of them). Finer than the quick-match vocabulary.
COVERAGE_EXTRA = [(r'azure functions', 'Azure Functions'), (r'data factory', 'Azure Data Factory'), (r'container apps', 'Azure Container Apps'),
                  (r'app service', 'Azure App Service'), (r'snowflake', 'Snowflake'), (r'databricks', 'Databricks'), (r'bigquery', 'BigQuery'),
                  (r'dynamo ?db', 'DynamoDB'), (r'\bsqs\b', 'SQS'), (r'\blambda\b', 'AWS Lambda'), (r'neo4j', 'Neo4j'), (r'json-ld', 'JSON-LD'),
                  (r'masstransit', 'MassTransit'), (r'entity framework|\bef core\b', 'Entity Framework'), (r'\bgrpc\b', 'gRPC'),
                  (r'event[- ]driven', 'event-driven'), (r'\bddd\b|domain[- ]driven', 'DDD')]

def cv_gaps(posting, cv_text):
    """Technologies the posting names that the CV text never mentions (same regex on both sides). Recruiters and ATS filters
    match words, so a skill the user has but the CV omits counts as missing."""
    import linkedin_jobs as lj
    out = []
    for rx, label in COVERAGE_EXTRA + lj.TECH:
        if re.search(rx, posting or '', re.I) and not re.search(rx, cv_text or '', re.I) and label not in out: out.append(label)
    return out

def _docx_text(p):
    import docx
    return '\n'.join(par.text for par in docx.Document(p).paragraphs)

def report_gaps(jid, v, docx_p):
    """Prints and saves (v['cv_gaps']) the posting's named technologies missing from the CV. LinkedIn jobs only."""
    if not str(jid).isdigit(): return
    try:
        import linkedin_jobs as lj
        gaps = cv_gaps(lj.detail(str(jid)).get('description') or '', _docx_text(docx_p))
    except Exception as e:
        print(f'(conferência de palavras do anúncio não rodou: {str(e)[:80]})'); return
    v['cv_gaps'] = gaps
    if gaps: print('O ANÚNCIO CITA E O CV NÃO: ' + ', '.join(gaps) + '\n   (se você tem a experiência, citar no Summary, na linha de Skills ou num bloco; se não tem, é lacuna real)')
    else: print('palavras do anúncio: o CV cita todas as tecnologias nomeadas')

def finalize(a, cfg):
    sp, s = _state(); v = (s.get('jobs') or {}).get('tracked', {}).get(str(a.id))
    if not v or not v.get('cv_docx'): sys.exit('rode primeiro a montagem (--profile) para esta vaga')
    pdf = os.path.splitext(v['cv_docx'])[0] + '.pdf'
    r = subprocess.run([sys.executable, os.path.join(HERE, 'cv_pdf.py'), a.pdf, pdf, v['cv_docx']], capture_output=True, text=True)
    print(r.stdout.strip() or r.stderr.strip())
    if 'texto igual ao DOCX: True' not in r.stdout: print('ATENCAO: o PDF nao bate com o DOCX; conferir antes de enviar')
    v.update(cv_pdf=pdf, cv_gdoc=a.gdoc.split('?')[0]); _save(sp, s)
    print(f'PDF pronto para anexar: {pdf}\nGoogle Doc: {v["cv_gdoc"]} (vai para a coluna Candidatura no Notion no proximo sync)')

def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('id', nargs='?'); ap.add_argument('--profile', '--perfil', dest='perfil'); ap.add_argument('--summary'); ap.add_argument('--skills')
    ap.add_argument('--company', '--empresa', dest='empresa'); ap.add_argument('--force', '--forcar', dest='forcar', action='store_true'); ap.add_argument('--profiles', '--perfis', dest='perfis', action='store_true')
    ap.add_argument('--pdf'); ap.add_argument('--gdoc'); ap.add_argument('--timeline', '--cronologia', dest='cronologia', metavar='DOCX')
    a = ap.parse_args(); cfg = criteria.load()
    if a.perfis:
        for n, p in (cfg.get('cv_profiles') or {}).items(): print(f'{n}: ' + ' > '.join(e[1].split(' at ')[-1] for e in p.get('experiences', [])[:6]) + ' …')
        return
    if a.cronologia:
        exps, prob = timeline(a.cronologia); print('\n'.join(prob) or 'cronologia ok'); return
    if not a.id: ap.error('informe o id da vaga')
    if a.pdf:
        if not a.gdoc: ap.error('--pdf pede --gdoc')
        return finalize(a, cfg)
    if not (a.perfil and a.summary): ap.error('informe --profile e --summary')
    assemble(a, cfg)

if __name__ == '__main__':
    main()
